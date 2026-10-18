"""File parsers for dealership-provided inventory: CSV/TSV and Excel (.xlsx)."""
from __future__ import annotations

import csv
import io
import json
from typing import Any

from app.core.errors import BadRequest


def parse_csv(data: bytes | str) -> tuple[list[str], list[dict[str, Any]]]:
    text = data.decode("utf-8-sig", errors="replace") if isinstance(data, bytes) else data.lstrip("﻿")
    if not text.strip():
        return [], []
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample.splitlines()[0], delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    try:
        rows = [r for r in reader if any(c.strip() for c in r)]
    except csv.Error as exc:
        raise BadRequest(f"Malformed CSV: {exc}") from exc
    if not rows:
        return [], []
    headers = [h.strip() for h in rows[0]]
    out = []
    for i, cells in enumerate(rows[1:], start=2):
        rec = {h: (cells[j].strip() if j < len(cells) else "") for j, h in enumerate(headers) if h}
        if len(cells) > len(headers) and any(c.strip() for c in cells[len(headers):]):
            rec["__extra_cells"] = cells[len(headers):]
        rec["__row_number"] = i
        out.append(rec)
    return headers, out


def parse_xlsx(data: bytes, sheet: str | None = None) -> tuple[list[str], list[dict[str, Any]]]:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise BadRequest(f"Could not read spreadsheet: {exc}") from exc
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = [["" if v is None else str(v).strip() for v in r] for r in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(rows and c for c in r)]
    if not rows:
        return [], []
    headers = rows[0]
    out = []
    for i, cells in enumerate(rows[1:], start=2):
        rec = {h: (cells[j] if j < len(cells) else "") for j, h in enumerate(headers) if h}
        rec["__row_number"] = i
        out.append(rec)
    return headers, out


def parse_json_feed(data: bytes | str, records_path: str | None = None) -> tuple[list[str], list[dict[str, Any]]]:
    try:
        doc = json.loads(data)
    except json.JSONDecodeError as exc:
        raise BadRequest(f"Malformed JSON feed: {exc}") from exc
    for key in (records_path or "").split(".") if records_path else []:
        doc = doc.get(key, []) if isinstance(doc, dict) else []
    if isinstance(doc, dict):
        doc = doc.get("vehicles") or doc.get("inventory") or doc.get("data") or []
    if not isinstance(doc, list):
        raise BadRequest("JSON feed must contain a list of vehicle records")
    headers: list[str] = []
    out = []
    for i, rec in enumerate(doc, start=1):
        if not isinstance(rec, dict):
            continue
        flat = {k: v for k, v in rec.items()}
        for k in flat:
            if k not in headers:
                headers.append(k)
        flat["__row_number"] = i
        out.append(flat)
    return headers, out


def parse_file(filename: str, data: bytes, content_type: str | None = None) -> tuple[str, list[str], list[dict[str, Any]]]:
    name = (filename or "").lower()
    if name.endswith(".xlsx") or (content_type or "").endswith("spreadsheetml.sheet"):
        return ("xlsx", *parse_xlsx(data))
    if name.endswith(".xls"):
        raise BadRequest("Legacy .xls files are not supported; save as .xlsx or CSV")
    if name.endswith(".json"):
        return ("json", *parse_json_feed(data))
    return ("csv", *parse_csv(data))
