"""CSV / XLSX import with column mapping, dry runs and row-level errors."""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from typing import Any, Iterable

from sqlalchemy.orm import Session

from ..audit import audit
from . import txdmv
from .prospects import LeadRecord, intake

# Accepted header spellings -> field
HEADER_ALIASES = {
    "name": "name", "dealership": "name", "dealership name": "name", "business name": "name", "company": "name",
    "website": "website", "url": "website", "web site": "website",
    "inventory url": "inventory_url", "inventory": "inventory_url", "inventory_url": "inventory_url",
    "address": "address_line", "street": "address_line", "address line": "address_line", "address_line": "address_line",
    "city": "city", "state": "state", "zip": "postal_code", "zip code": "postal_code", "postal code": "postal_code",
    "postal_code": "postal_code", "phone": "phone", "phone number": "phone", "business phone": "phone",
    "email": "email", "business email": "email", "type": "dealership_type", "dealership type": "dealership_type",
    "dealership_type": "dealership_type", "notes": "notes", "source url": "source_url", "source_url": "source_url",
    "latitude": "lat", "lat": "lat", "longitude": "lng", "lng": "lng", "lon": "lng",
    # Spellings seen in licensing lists (e.g. the TxDMV dealer list); check your file's headers in the dry run.
    "license holder": "name", "licensee name": "name", "legal name": "name",
    "physical address": "address_line", "street address": "address_line", "address 1": "address_line",
    "physical city": "city", "physical state": "state", "physical zip": "postal_code", "zip code 5": "postal_code",
    "telephone": "phone", "phone #": "phone",
    "license number": "license_number", "license #": "license_number", "gdn": "license_number",
    "license type": "license_type", "license status": "license_status", "county": "county",
    "dba": "dba", "assumed name": "dba", "assumed name/dba": "dba", "doing business as": "dba",
}
# Columns stored as evidence (not prospect fields).
EVIDENCE_FIELDS = {"license_number": "license:number", "license_type": "license:type",
                   "license_status": "license:status", "county": "county", "dba": "dba"}
MAX_ROWS = 5000


@dataclass
class ImportReport:
    total: int = 0
    created: int = 0
    merged: int = 0
    skipped: int = 0
    dry_run: bool = False
    rows: list[dict[str, Any]] = field(default_factory=list)
    excluded: dict[str, int] = field(default_factory=dict)  # rows filtered out before intake, by reason


def read_rows(filename: str, content: bytes) -> list[dict[str, str]]:
    if filename.lower().endswith(".xls") or content[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        # Legacy Excel 97-2003 (the TxDMV dealer list is published in this format)
        import xlrd

        sh = xlrd.open_workbook(file_contents=content).sheet_by_index(0)
        headers = [" ".join(str(h).split()) for h in sh.row_values(0)]
        rows = []
        for r in range(1, sh.nrows):
            vals = sh.row_values(r)
            rows.append({headers[i]: _cell(v) for i, v in enumerate(vals) if i < len(headers)})
        return [r for r in rows if any(v.strip() for v in r.values())]
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        headers = [str(h or "").strip() for h in next(it, [])]
        rows = [{headers[i]: ("" if v is None else str(v)) for i, v in enumerate(r) if i < len(headers)} for r in it]
        return [r for r in rows if any(v.strip() for v in r.values())]
    text = content.decode("utf-8-sig", errors="replace")
    return [r for r in csv.DictReader(io.StringIO(text)) if any((v or "").strip() for v in r.values())]


def _cell(v: Any) -> str:
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return "" if v is None else str(v)


def map_headers(headers: Iterable[str]) -> dict[str, str]:
    return {h: HEADER_ALIASES[h.strip().lower()] for h in headers if h and h.strip().lower() in HEADER_ALIASES}


def import_rows(db: Session, rows: list[dict[str, str]], *, source_label: str, source_url: str | None = None,
                dry_run: bool = False, actor_id: int | None = None, simulated: bool = False) -> ImportReport:
    """Import rows. ``source_label`` names where the file came from (shown as the lead source)."""
    report = ImportReport(total=len(rows), dry_run=dry_run)
    if rows and txdmv.is_txdmv(rows[0].keys()):
        raise ValueError("This is a TxDMV dealer list; import it with the TxDMV option (filters by county/city)")
    if len(rows) > MAX_ROWS:
        raise ValueError(f"File has {len(rows)} rows; the limit per import is {MAX_ROWS}")
    mapping = map_headers(rows[0].keys()) if rows else {}
    if rows and not ({"name", "dba"} & set(mapping.values())):
        raise ValueError("No dealership name column found (expected a header like 'name' or 'dealership name')")
    records = []
    for i, row in enumerate(rows, start=2):  # line 1 is the header
        vals: dict[str, Any] = {}
        for header, fld in mapping.items():
            v = (row.get(header) or "").strip()
            if v:
                vals[fld] = v
        if vals.get("dealership_type"):
            vals["dealership_type"] = vals["dealership_type"].lower().replace(" ", "_").replace("-", "_")
        signals = [(EVIDENCE_FIELDS[k], vals.pop(k), f"From imported list '{source_label}'")
                   for k in list(vals) if k in EVIDENCE_FIELDS]
        if not vals.get("name") and any(f == "dba" for f, _, _ in signals):
            vals["name"] = next(v for f, v, _ in signals if f == "dba")
        records.append((i, LeadRecord(
            name=vals.get("name", ""), source=f"csv:{source_label}"[:60],
            source_url=vals.pop("source_url", None) or source_url, simulated=simulated,
            notes=vals.pop("notes", ""), signals=signals, **{k: v for k, v in vals.items() if k not in ("name",)},
        )))
    return intake_records(db, records, report, source_label=source_label, actor_id=actor_id)


def intake_records(db: Session, records: list[tuple[int, LeadRecord]], report: ImportReport, *, source_label: str,
                   actor_id: int | None) -> ImportReport:
    savepoint = db.begin_nested() if report.dry_run else None
    for i, rec in records:
        res = intake(db, rec, actor_id=actor_id)
        if res.prospect is None:
            report.skipped += 1
            status = "skipped"
        elif res.created:
            report.created += 1
            status = "created"
        else:
            report.merged += 1
            status = f"merged ({res.matched_by})"
        report.rows.append({"line": i, "name": rec.name, "status": status,
                            "prospect_id": res.prospect.id if res.prospect and not report.dry_run else None,
                            "errors": res.errors, "conflicts": res.conflicts})
    if savepoint is not None:
        savepoint.rollback()
    else:
        audit(db, "import.file", actor_id=actor_id, data={"source": source_label, "created": report.created,
                                                           "merged": report.merged, "skipped": report.skipped,
                                                           "excluded": report.excluded})
    return report


def import_txdmv(db: Session, rows: list[dict[str, str]], *, counties: list[str], cities: list[str],
                 include_franchise: bool = False, data_date: str = "", dry_run: bool = False,
                 actor_id: int | None = None) -> ImportReport:
    if not rows or not txdmv.is_txdmv(rows[0].keys()):
        raise ValueError("This file doesn't have the TxDMV dealer list columns (LicenseNumber, DBAName, ...)")
    records, excluded = txdmv.to_records(rows, counties=counties, cities=cities,
                                         include_franchise=include_franchise, data_date=data_date)
    if len(records) > MAX_ROWS:
        raise ValueError(f"{len(records)} dealers match; narrow the counties/cities (limit {MAX_ROWS} per import)")
    report = ImportReport(total=len(rows), dry_run=dry_run, excluded=dict(excluded))
    return intake_records(db, records, report, source_label="TxDMV dealer list", actor_id=actor_id)
